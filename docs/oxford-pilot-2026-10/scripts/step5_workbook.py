# Step 5: write the lead-check workbook.
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

out = sys.argv[1]
s = pd.read_csv("sites.csv").fillna("")
b = pd.read_csv("buildings_final.csv").fillna("")

s["rank_all_sites"] = range(1, len(s) + 1)
leads = s[s["repd_same_site"] == ""].reset_index(drop=True)
leads["rank"] = range(1, len(leads) + 1)
leads["score_0_10"] = (10 * leads["points_best"] / 82).round(1)

def flags(r):
    f = []
    if r["pv_elsewhere_on_site"]:
        f.append("Panels already mapped on %d other building(s) here" % r["pv_elsewhere_on_site"])
    if r["repd_nearby"]:
        f.append("Registered solar project nearby: " + r["repd_nearby"].split("|")[0].strip())
    if r["heritage_osm"]:
        f.append("Heritage tag in OSM: " + r["heritage_osm"])
    return "; ".join(f)
leads["flags"] = leads.apply(flags, axis=1)

def sat_link(lat, lon):
    return "https://www.google.com/maps/@?api=1&map_action=map&center=%.6f,%.6f&zoom=19&basemap=satellite" % (lat, lon)
def osm_link(oid):
    return "https://www.openstreetmap.org/%s/%s" % ("way" if oid.startswith("w") else "relation", oid[1:])

HEAD = PatternFill("solid", fgColor="DDE4EE")
FILL_ME = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")

CHECK_COLS = [("Panels visible?", ["Yes", "No", "Unsure"]),
              ("Roof usable?", ["Yes", "Partly", "No"]),
              ("Occupier (confirmed)", None),
              ("Same site as rank #", None),
              ("Keep for Chris?", ["Yes", "No", "Maybe"]),
              ("Notes", None)]

def lead_sheet(wb, title, df, extra_cols=()):
    ws = wb.create_sheet(title)
    cols = [("Rank", "rank", 7), ("Site (best guess from OSM)", "site", 40), ("Sector (OSM)", "sector", 24),
            ("Buildings", "buildings", 10), ("Largest roof m²", "largest_roof_m2", 15), ("Total roof m²", "total_roof_m2", 14),
            ("Score 0-10", "score_0_10", 11), ("Site gap m", "site_gap_m", 11)] + list(extra_cols) + \
           [("Address (OSM)", "address_osm", 28), ("Flags", "flags", 40)]
    headers = [c[0] for c in cols] + ["Satellite view", "OSM"] + [c[0] for c in CHECK_COLS]
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = BOLD
        cell.fill = FILL_ME if i > len(cols) + 2 else HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for _, r in df.iterrows():
        ws.append([r[c[1]] for c in cols] + ["Satellite", "OSM"] + [""] * len(CHECK_COLS))
        row = ws.max_row
        c1 = ws.cell(row=row, column=len(cols) + 1); c1.hyperlink = sat_link(r["lat"], r["lon"]); c1.font = LINK
        c2 = ws.cell(row=row, column=len(cols) + 2); c2.hyperlink = osm_link(r["lead_oid"]); c2.font = LINK
    widths = [c[2] for c in cols] + [11, 7, 14, 13, 26, 14, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for j, (name, options) in enumerate(CHECK_COLS):
        col = get_column_letter(len(cols) + 3 + j)
        if options:
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("%s2:%s%d" % (col, col, ws.max_row))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws

wb = Workbook()
# --- Sheet 1: how to use it -------------------------------------------------------------------
ws = wb.active
ws.title = "How to check"
n_pv = int(b["solar_now"].sum())
lines = [
    ("Oxford City lead check, 7 October 2026", True),
    ("", False),
    ("What this is", True),
    ("Commercial and industrial sites in Oxford City (local authority E07000178), ranked with the OZ scorecard v0.3 (pilot): roof footprint area plus site gap, the method that won the arena run.", False),
    ("Buildings come from OpenStreetMap. Buildings on one site with one owner (a school, hospital, college or single-company site) are grouped into one lead.", False),
    ("", False),
    ("Already done automatically", True),
    ("1,790 buildings scored; %d removed because OpenStreetMap shows solar panels on them." % n_pv, False),
    ("%d sites removed because they already have a rooftop solar planning record in the REPD (see the 'Already in solar pipeline' tab)." % int((s.repd_same_site != '').sum()), False),
    ("%d sites remain. Tabs: Top 50 (by score), Large roofs (largest building 2,375 m² or more, ranked by size, as the method recommends for 150 kWp+ projects), Reserves (ranks 51-150)." % len(leads), False),
    ("", False),
    ("How to check each lead (about 2 minutes each; do the Top 50 first, then any Large roofs not already in it)", True),
    ("1. Click 'Satellite'. It opens Google Maps satellite view on the building.", False),
    ("2. Panels visible? Yes / No / Unsure. Imagery can be one to three years old, so 'Unsure' is fine.", False),
    ("3. Roof usable? Yes / Partly (lots of plant, skylights, odd shape, shading) / No.", False),
    ("4. Occupier: use the map labels or Street View and write the business name. Rows called '(unit on site)' or 'Unnamed' need this most.", False),
    ("5. Same site as rank #: if two rows are the same business, write the other rank so we merge them.", False),
    ("6. Keep for Chris? Yes for the best ~20: no panels, usable roof, known occupier, a mix of sectors.", False),
    ("7. For your final 20 only: check Historic England's map for listed buildings and conservation areas. This could not be done automatically (the planning data service is blocked on this network).", False),
    ("", False),
    ("Not done yet", True),
    ("The Feed-in Tariff 'likely unmapped panels' check (method step 7b); the official heritage check; kWp, savings and CO2 estimates (next step, after this check).", False),
    ("", False),
    ("Read the numbers carefully", True),
    ("OpenStreetMap records only part of the solar already installed, so expect to find panels the data missed. Site names are best guesses from OSM tags.", False),
    ("The score ranks sites against each other; it is not a probability of buying. Pilot release: on 150 kWp+ projects the scorecard did worse than roof size alone, hence the Large roofs tab.", False),
    ("Of the 11 Oxford sites with their own rooftop solar planning record, 8 ranked in the top 80 of %d sites before they were removed (ranks: %s). An encouraging sign, not proof." % (len(s), ", ".join(str(x) for x in s[s.repd_same_site != ''].rank_all_sites)), False),
    ("", False),
    ("Sources and attribution", True),
    ("Contains OpenStreetMap data © OpenStreetMap contributors, available under the Open Database Licence (ODbL).", False),
    ("Contains public sector information licensed under the Open Government Licence v3.0 (Renewable Energy Planning Database, DESNZ).", False),
    ("Scores: OZ arena run 2 Oct 2026, champion a006, scorecard v0.3. Satellite links open Google Maps for viewing only; nothing from the imagery is stored except your own notes.", False),
]
for text, bold in lines:
    ws.append([text])
    if bold:
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13 if ws.max_row == 1 else 11)
ws.column_dimensions["A"].width = 140
for row in ws.iter_rows():
    row[0].alignment = Alignment(wrap_text=True, vertical="top")

# --- Sheets 2-4: the lists -------------------------------------------------------------------
lead_sheet(wb, "Top 50", leads.head(50))
big = leads[leads["largest_roof_m2"] >= 2375].sort_values(["largest_roof_m2", "lead_oid"], ascending=[False, True]).head(25).copy()
big["in_top50"] = [("yes, #%d" % r) if r <= 50 else "" for r in big["rank"]]
lead_sheet(wb, "Large roofs", big, extra_cols=[("Also in Top 50?", "in_top50", 12)])
lead_sheet(wb, "Reserves", leads.iloc[50:150])

# --- Sheet 5: already in the solar pipeline -----------------------------------------------------
ws = wb.create_sheet("Already in solar pipeline")
ws.append(["Rank among all sites", "Site", "Sector", "Largest roof m²", "Score 0-10", "REPD record (name | MW | status [distance from building])"])
for _, r in s[s.repd_same_site != ""].iterrows():
    ws.append([r["rank_all_sites"], r["site"], r["sector"], r["largest_roof_m2"], round(10 * r["points_best"] / 82, 1), r["repd_same_site"]])
for i, w in enumerate([12, 42, 22, 15, 11, 110], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = BOLD; c.fill = HEAD

# --- Sheet 6: buildings excluded for mapped panels ---------------------------------------------
ws = wb.create_sheet("Excluded - panels mapped")
ws.append(["OSM id", "Building name (OSM)", "Site (OSM)", "Sector", "Roof m²", "Satellite view"])
for _, r in b[b["solar_now"]].sort_values("area", ascending=False).iterrows():
    ws.append([r["oid"], r["osm_name"], r["site_name"], r["sector"], round(r["area"]), "Satellite"])
    c = ws.cell(row=ws.max_row, column=6); c.hyperlink = sat_link(r["lat"], r["lon"]); c.font = LINK
for i, w in enumerate([14, 36, 36, 26, 10, 12], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = BOLD; c.fill = HEAD

wb.save(out)
print("saved:", out)
print("Top 50 rows:", min(50, len(leads)), "| Large roofs rows:", len(big), "(", int((big['rank'] <= 50).sum()), "also in Top 50 )",
      "| Reserves:", len(leads.iloc[50:150]), "| pipeline:", int((s.repd_same_site != '').sum()), "| excluded buildings:", int(b['solar_now'].sum()))
